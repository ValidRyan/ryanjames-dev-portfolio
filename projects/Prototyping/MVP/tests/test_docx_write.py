"""Round-trip tests for docx_write.py against File Recompiler's docx_extract.

Uses the exact sys.path.insert(...) convention from test_docx_extract.py to
reach both File Recompiler/ (for docx_extract) and Joint Interactor/ (for docx_write).
"""

import sys
from pathlib import Path

_MVP_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_MVP_ROOT / "File Recompiler"))
sys.path.insert(0, str(_MVP_ROOT / "Joint Interactor"))

import pytest

import docx_extract
from docx_write import write_docx


@pytest.mark.parametrize(
    "text",
    [
        "",
        "single line",
        "line one\nline two\nline three",
        "trailing newline\n",
        "leading empty line\n\nsecond paragraph",
    ],
)
def test_round_trip_extract_text_matches_original(tmp_path, text):
    target = tmp_path / "roundtrip.docx"
    write_docx(text, str(target))

    result = docx_extract.extract_text(str(target))

    assert result == text


def test_write_docx_creates_one_paragraph_per_line(tmp_path):
    target = tmp_path / "multiline.docx"
    write_docx("a\nb\nc", str(target))

    from docx import Document

    document = Document(str(target))
    assert [p.text for p in document.paragraphs] == ["a", "b", "c"]


def test_write_docx_empty_text_creates_single_empty_paragraph(tmp_path):
    target = tmp_path / "empty.docx"
    write_docx("", str(target))

    from docx import Document

    document = Document(str(target))
    assert [p.text for p in document.paragraphs] == [""]


def test_write_docx_accepts_path_object(tmp_path):
    target = tmp_path / "path_object.docx"
    write_docx("hello", target)

    assert docx_extract.extract_text(str(target)) == "hello"
